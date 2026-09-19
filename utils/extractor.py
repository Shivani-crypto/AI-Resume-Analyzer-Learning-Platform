import os
import re
import shutil
from typing import Optional

import pdfplumber
import docx
import pytesseract
from PIL import Image
import pypdf

# -----------------------------
# Dynamic Tesseract Path Detection
# -----------------------------
def get_tesseract_cmd() -> Optional[str]:
    cmd = shutil.which("tesseract")
    if cmd and os.path.exists(cmd):
        return cmd
    possible_paths = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
        os.path.expanduser(r"~\AppData\Local\Tesseract-OCR\tesseract.exe")
    ]
    for p in possible_paths:
        if os.path.exists(p):
            return p
    return None

tesseract_path = get_tesseract_cmd()
if tesseract_path:
    pytesseract.pytesseract.tesseract_cmd = tesseract_path


# -----------------------------
# Helper: PDF Text Extraction
# -----------------------------
def extract_text_from_pdf(filepath: str) -> str:
    text = ""
    # 1. Primary Method: pdfplumber
    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                try:
                    content = page.extract_text()
                    if not content or not content.strip():
                        content = page.extract_text(x_tolerance=2, y_tolerance=2)
                    if content and content.strip():
                        text += content + "\n"
                except Exception as pe:
                    print(f"[Extractor] pdfplumber page extraction error: {pe}")
    except Exception as e:
        print(f"[Extractor] pdfplumber open error: {e}")

    # 2. Fallback Method: pypdf (if pdfplumber extracted nothing or very little)
    if len(text.strip()) < 30:
        try:
            reader = pypdf.PdfReader(filepath)
            pypdf_text = ""
            for page in reader.pages:
                t = page.extract_text()
                if t and t.strip():
                    pypdf_text += t + "\n"
            if len(pypdf_text.strip()) > len(text.strip()):
                text = pypdf_text
        except Exception as e:
            print(f"[Extractor] pypdf extraction error: {e}")

    # 3. OCR Fallback for scanned PDFs (if tesseract is installed and text is still empty)
    if len(text.strip()) < 30 and tesseract_path:
        try:
            with pdfplumber.open(filepath) as pdf:
                ocr_text = ""
                for page in pdf.pages:
                    try:
                        img_obj = page.to_image(resolution=200).original
                        page_ocr = pytesseract.image_to_string(img_obj)
                        if page_ocr and page_ocr.strip():
                            ocr_text += page_ocr + "\n"
                    except Exception as ocr_err:
                        print(f"[Extractor] PDF page OCR error: {ocr_err}")
                if len(ocr_text.strip()) > len(text.strip()):
                    text = ocr_text
        except Exception as e:
            print(f"[Extractor] PDF OCR fallback error: {e}")

    return text.strip()


# -----------------------------
# Helper: DOCX & DOC Extraction
# -----------------------------
def extract_text_from_docx(filepath: str) -> str:
    text = ""
    try:
        doc = docx.Document(filepath)
        # 1. Extract paragraphs
        for para in doc.paragraphs:
            if para.text and para.text.strip():
                text += para.text + "\n"

        # 2. Extract tables (Resumes often put skills/experience inside table cells!)
        for table in doc.tables:
            for row in table.rows:
                cell_texts = []
                for cell in row.cells:
                    ctext = cell.text.strip()
                    if ctext and ctext not in cell_texts:
                        cell_texts.append(ctext)
                if cell_texts:
                    text += " | ".join(cell_texts) + "\n"
    except Exception as e:
        print(f"[Extractor] docx extraction error: {e}")

    # 3. Fallback for binary .doc or malformed docx files
    if not text.strip():
        try:
            with open(filepath, "rb") as f:
                content = f.read()
                matches = re.findall(rb'[\x20-\x7E\t\r\n]{4,}', content)
                fallback_strings = [m.decode('ascii', errors='ignore').strip() for m in matches]
                cleaned_strings = [s for s in fallback_strings if not s.startswith('<') and len(s) > 3]
                if cleaned_strings:
                    text = "\n".join(cleaned_strings)
        except Exception as fe:
            print(f"[Extractor] binary file string fallback error: {fe}")

    return text.strip()


# -----------------------------
# Helper: Image OCR Extraction
# -----------------------------
def extract_text_from_image(filepath: str) -> str:
    if not tesseract_path:
        print("[Extractor] Tesseract OCR executable not found.")
        return ""

    try:
        img = Image.open(filepath)
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        text = pytesseract.image_to_string(img)
        return text.strip()
    except Exception as e:
        print(f"[Extractor] Image OCR error: {e}")
        return ""


# -----------------------------
# 1. EXTRACT RAW TEXT
# -----------------------------
def extract_raw_text(filepath: str) -> str:
    if not os.path.exists(filepath):
        return ""

    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(filepath)
    elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tiff"]:
        return extract_text_from_image(filepath)
    else:
        return ""


# -----------------------------
# 2. CLEAN TEXT
# -----------------------------
def clean_text(text: str) -> str:
    if not text:
        return ""

    # Fix broken email domains (e.g., user@ gmail. com -> user@gmail.com)
    text = re.sub(r'(\b[A-Za-z0-9._%+-]+)\s*@\s*([A-Za-z0-9.-]+)\s*\.\s*([A-Za-z]{2,}\b)', r'\1@\2.\3', text)

    # Normalize line breaks while preserving document layout
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    cleaned_lines = [re.sub(r'[ \t]+', ' ', line) for line in lines]

    return "\n".join(cleaned_lines)


# -----------------------------
# 3. FORMAT INTO RESUME STYLE
# -----------------------------
def format_resume(text: str) -> str:
    if not text:
        return ""

    sections = [
        "PROFESSIONAL SUMMARY",
        "WORK EXPERIENCE",
        "EXPERIENCE",
        "EDUCATION",
        "TECHNICAL SKILLS",
        "SKILLS",
        "PROJECTS",
        "CERTIFICATIONS"
    ]

    for sec in sections:
        text = re.sub(rf'\b({sec})\b', r'\n\n\1\n', text, flags=re.IGNORECASE)

    text = text.replace("•", "\n•")
    text = re.sub(r'\n{3,}', '\n\n', text)

    return text.strip()


# -----------------------------
# 4. MAIN FUNCTION
# -----------------------------
def extract_text_from_file(filepath: str) -> str:
    raw_text = extract_raw_text(filepath)
    if not raw_text or not raw_text.strip():
        return ""

    cleaned = clean_text(raw_text)
    formatted = format_resume(cleaned)
    return formatted


# -----------------------------
# 5. TEST (Run directly)
# -----------------------------
if __name__ == "__main__":
    file_path = "sample_resume.pdf"
    result = extract_text_from_file(file_path)
    print("\n===== CLEAN FORMATTED RESUME =====\n")
    print(result)